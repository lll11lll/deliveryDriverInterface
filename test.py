# Nick Greiner
# 01/15/2023
# Testing myDriverData.xlsx

# I am testing the tables and how to extend the table ref by one for my program

# I just want to see if it will identitfy the tables
# i would to it after i insert the data then extend it by 1
import openpyxl as op
wb = op.load_workbook('testingData.xlsx')
sheet = wb['mainSheet']

for table in sheet.tables.items():
    print(table)

print(sheet.tables['Table2'].ref)
# print(sheet.tables['Table3'])

print(sheet.tables['Table2'].ref)

wb.save('testingData2.xlsx')


# Used to fix my program for wage amount, hourly rate and daily total
tempBugFix = 3
while tempBugFix != 9:
    sheet[f'F{tempBugFix}'].value = f'=C{tempBugFix} * 7.98'
    sheet[f'H{tempBugFix}'].value = f'=( F{tempBugFix} + B{tempBugFix} ) / C{tempBugFix}'
    sheet[f'I{tempBugFix}'].value = f'=B{tempBugFix} + F{tempBugFix}'
    tempBugFix += 1
